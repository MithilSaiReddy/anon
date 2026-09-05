import io
import json
import logging
import os
import re
import sys
import shutil
import time
import zipfile
from pathlib import Path
from typing import Dict, Any, Optional

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import FileResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware

#this is test to use fastapi-mcp library and also pydantic
from pydantic import BaseModel
from fastapi_mcp import FastApiMCP


import openpyxl
from docx import Document

from backend.anonymizer import create_anonymizer, Anonymizer
from backend.restorer import create_restorer, Restorer
from backend.file_handlers import (
    extract_text_from_docx, extract_text_from_xlsx, extract_text_from_pdf, extract_text_from_pptx,
    anonymize_docx, anonymize_xlsx, anonymize_pdf, anonymize_pptx,
    restore_docx, restore_xlsx, restore_pdf, restore_pptx,
)
from backend.logger import setup_logging
from backend.pdf_pipeline import is_scanned_pdf, process_scanned_pdf, cleanup_ocr_temp
from backend.memory_monitor import MemoryThresholdExceeded

logger = logging.getLogger(__name__)

app_env = os.getenv("APP_ENV", "dev")
setup_logging(app_env)

app = FastAPI()

#init mcp server 
mcp = FastApiMCP(app)
mcp.mount()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.exception_handler(MemoryThresholdExceeded)
async def memory_threshold_handler(request, exc):
    logger.warning("Memory threshold exceeded: %s", exc)
    return JSONResponse(
        status_code=413,
        content={"success": False, "detail": str(exc)},
    )


BASE_DIR = Path(os.path.abspath(__file__)).parent.parent
TEMP_DIR = BASE_DIR / "temp"
TEMP_DIR.mkdir(exist_ok=True)

anonymizer: Optional[Anonymizer] = None
restorer: Optional[Restorer] = None


@app.on_event("startup")
async def startup():
    global anonymizer, restorer
    logger.info("Starting Anon server (env=%s) ...", app_env)
    anonymizer = create_anonymizer()
    restorer = create_restorer()
    anonymizer.load_gliner()
    anonymizer.load_spacy()
    logger.info("Server ready!")


@app.get("/")
async def root():
    index_path = BASE_DIR / "frontend" / "index.html"
    return FileResponse(index_path, headers={"Cache-Control": "no-store"})


@app.get("/health")
async def health():
    return {"status": "ok", "gliner_loaded": anonymizer is not None and anonymizer.nlp is not None, "spacy_loaded": anonymizer is not None and anonymizer.nlp_spacy is not None}


# ── MCP TEXT TOOLS ──────────────────────────────────────
class TextIn(BaseModel):
    text: str

class AnonymizeTextRequest(BaseModel):
    text: str
    anonymize_persons: bool = True
    anonymize_orgs: bool = True
    anonymize_locations: bool = True
    anonymize_ids: bool = True

class RestoreTextRequest(BaseModel):
    anonymized_text: str
    bridge_key: dict


@app.post("/api/detect-text",
    summary="Preview how many sensitive entities exist in text",
    description="Scans text locally, returns ONLY counts of persons/orgs/locations/IDs. "
                "Never returns actual names. Call before anonymizing so user can confirm.")
async def detect_text(req: TextIn):
    if not req.text.strip():
        raise HTTPException(status_code=400, detail="Text cannot be empty")
    entity_map = anonymizer.extract_entities(req.text)
    counts = {
        "persons":   sum(1 for k in entity_map if k.startswith("PERSON_")),
        "orgs":      sum(1 for k in entity_map if k.startswith("ORG_")),
        "locations": sum(1 for k in entity_map if k.startswith("LOC_")),
        "ids":       sum(1 for k in entity_map if k.startswith("ID_")),
    }
    total = sum(counts.values())
    return {"counts": counts, "total": total,
            "message": f"Found {total} sensitive entities. Safe to anonymize."}


@app.post("/api/anonymize-text",
    summary="Anonymize sensitive text before sending to AI",
    description="Replaces names, orgs, locations, IDs with placeholders like PERSON_1, ORG_1. "
                "Returns anonymized_text (safe to send to LLM) and bridge_key (store this). "
                "Never send the bridge_key to the LLM.")
async def anonymize_text(req: AnonymizeTextRequest):
    if not req.text.strip():
        raise HTTPException(status_code=400, detail="Text cannot be empty")
    entity_types = {
        "PERSON": req.anonymize_persons,
        "ORG":    req.anonymize_orgs,
        "LOC":    req.anonymize_locations,
        "ID":     req.anonymize_ids,
    }
    entity_map, anonymized_text, stats = anonymizer.anonymize(req.text, 1.0, entity_types)
    bridge_key = anonymizer.create_bridge_key("text_input", entity_map, multiplier=1.0)
    return {
        "anonymized_text": anonymized_text,
        "bridge_key": bridge_key,
        "stats": {
            "persons":   stats.get("persons", 0),
            "orgs":      stats.get("orgs", 0),
            "locations": stats.get("locs", 0),
            "ids":       stats.get("ids", 0),
        },
    }


@app.post("/api/restore-text",
    summary="Restore anonymized text back to original using bridge key",
    description="Swaps PERSON_1, ORG_1 etc back to real values using the bridge_key "
                "from /api/anonymize-text. Call this on the LLM response to get real names back.")
async def restore_text(req: RestoreTextRequest):
    if not req.anonymized_text.strip():
        raise HTTPException(status_code=400, detail="Text cannot be empty")
    if not req.bridge_key:
        raise HTTPException(status_code=400, detail="Bridge key cannot be empty")
    restored = req.anonymized_text
    entities = req.bridge_key.get("entities", {})
    for placeholder, original in entities.items():
        restored = restored.replace(placeholder, original)
    replaced = sum(1 for p in entities if p in req.anonymized_text)
    return {"restored_text": restored, "replacements_made": replaced}


@app.post("/api/validate-key",
    summary="Validate a bridge key before using it",
    description="Checks if a bridge_key is valid and returns a count summary. "
                "Call before restore-text if unsure whether the key is intact.")
async def validate_key(bridge_key: dict):
    entities = bridge_key.get("entities", {})
    if not entities:
        return {"valid": False, "reason": "No entities found in bridge key"}
    counts = {
        "persons":   sum(1 for k in entities if k.startswith("PERSON_")),
        "orgs":      sum(1 for k in entities if k.startswith("ORG_")),
        "locations": sum(1 for k in entities if k.startswith("LOC_")),
        "ids":       sum(1 for k in entities if k.startswith("ID_")),
    }
    return {"valid": True, "total_entities": len(entities), "counts": counts,
            "original_file": bridge_key.get("original_filename", "text_input")}

@app.post("/api/detect")
async def detect_entities(file: UploadFile = File(...)):
    if not file.filename:
        raise HTTPException(status_code=400, detail="No file provided")

    original_filename = file.filename
    file_ext = os.path.splitext(original_filename)[1].lower()
    file_bytes = await file.read()

    try:
        logger.info("Detect: %s (%d bytes)", original_filename, len(file_bytes))
        if file_ext == ".docx":
            text, _ = extract_text_from_docx(file_bytes)
        elif file_ext == ".xlsx":
            text, _ = extract_text_from_xlsx(file_bytes)
        elif file_ext == ".pdf":
            if is_scanned_pdf(file_bytes):
                text, _ = process_scanned_pdf(file_bytes, TEMP_DIR)
            else:
                text, _ = extract_text_from_pdf(file_bytes)
        elif file_ext == ".pptx":
            text, _ = extract_text_from_pptx(file_bytes)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported file type: {file_ext}")

        entity_map = anonymizer.extract_entities(text)

        counts = {
            "PERSON": sum(1 for k in entity_map if k.startswith("PERSON_")),
            "ORG": sum(1 for k in entity_map if k.startswith("ORG_")),
            "LOC": sum(1 for k in entity_map if k.startswith("LOC_")),
            "ID": sum(1 for k in entity_map if k.startswith("ID_")),
        }

        return JSONResponse({
            "success": True,
            "entities": counts,
            "filename": original_filename,
            "entity_map": entity_map,
            "preview": {k: v for k, v in list(entity_map.items())[:10]}
        })

    except MemoryThresholdExceeded:
        logger.warning("Memory threshold exceeded during detect")
        raise HTTPException(status_code=413, detail="File Too Large — processing exceeded memory limit")
    except Exception as e:
        logger.exception("Detect failed: %s", e)
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/anonymize")
async def anonymize_file(
    file: UploadFile = File(...),
    multiplier: float = Form(1.0),
    anonymize_person: bool = Form(True),
    anonymize_org: bool = Form(True),
    anonymize_city: bool = Form(True),
    anonymize_id: bool = Form(True),
    skip_entities: Optional[str] = Form(None),
):
    if not file.filename:
        raise HTTPException(status_code=400, detail="No file provided")

    skip_set = set()
    if skip_entities:
        try:
            skip_set = set(json.loads(skip_entities))
        except (json.JSONDecodeError, TypeError):
            skip_set = {s_ for s_ in skip_entities.split(",") if s_.strip()}

    original_filename = file.filename
    file_ext = os.path.splitext(original_filename)[1].lower()
    file_bytes = await file.read()

    entity_types = {
        "PERSON": anonymize_person,
        "ORG": anonymize_org,
        "LOC": anonymize_city,
        "ID": anonymize_id,
    }

    try:
        start_time = time.time()
        logger.info("Processing: %s (ext=%s, mult=%.2f)", original_filename, file_ext, multiplier)

        if file_ext == ".docx":
            text, doc = extract_text_from_docx(file_bytes)
            entity_map, _, stats = anonymizer.anonymize(text, multiplier, entity_types, skip_entities=skip_set)

            counts, number_map = anonymize_docx(doc, entity_map, multiplier)

            output_buffer = io.BytesIO()
            doc.save(output_buffer)
            output_buffer.seek(0)
            anon_filename = f"anon_{original_filename}"
            anon_path = TEMP_DIR / anon_filename
            with open(anon_path, "wb") as f:
                f.write(output_buffer.getvalue())

            key_filename = f"{os.path.splitext(original_filename)[0]}.bridgekey.json"
            key_data = anonymizer.create_bridge_key(original_filename, entity_map, multiplier)
            key_data["number_mappings"] = number_map
            key_path = TEMP_DIR / key_filename
            with open(key_path, "w") as f:
                json.dump(key_data, f, indent=2)

            total_entities = stats["persons"] + stats["orgs"] + stats.get("locs", 0)
            elapsed = time.time() - start_time
            logger.info("Done: %s — %d entities, %d numbers (%.2fs)", original_filename, total_entities, counts['numbers'], elapsed)

            message = f"Anonymized {total_entities} entities, {counts['numbers']} numbers"
            return JSONResponse({
                "success": True,
                "message": message,
                "anon_filename": anon_filename,
                "key_filename": key_filename,
                "entity_mapping": entity_map,
                "stats": stats,
            })

        elif file_ext == ".xlsx":
            text, wb = extract_text_from_xlsx(file_bytes)
            logger.debug("Extracted %d chars from XLSX", len(text))

            entity_map, _, stats = anonymizer.anonymize(text, multiplier, entity_types, skip_entities=skip_set)

            counts, number_map = anonymize_xlsx(wb, entity_map, multiplier)

            anon_filename = f"anon_{original_filename}"
            anon_path = TEMP_DIR / anon_filename
            wb.save(anon_path)

            key_filename = f"{os.path.splitext(original_filename)[0]}.bridgekey.json"
            key_data = anonymizer.create_bridge_key(original_filename, entity_map, multiplier)
            key_data["number_mappings"] = number_map
            key_path = TEMP_DIR / key_filename
            with open(key_path, "w") as f:
                json.dump(key_data, f, indent=2)

            total_entities = stats["persons"] + stats["orgs"] + stats.get("locs", 0)
            elapsed = time.time() - start_time
            logger.info("Done: %s — %d entities, %d numbers (%.2fs)", original_filename, total_entities, counts['numbers'], elapsed)

            message = f"Anonymized {total_entities} entities, {counts['numbers']} numbers"
            return JSONResponse({
                "success": True,
                "message": message,
                "anon_filename": anon_filename,
                "key_filename": key_filename,
                "entity_mapping": entity_map,
                "stats": stats,
            })

        elif file_ext == ".pdf":
            if is_scanned_pdf(file_bytes):
                markdown_text, md_path = process_scanned_pdf(file_bytes, TEMP_DIR)
                entity_map, anon_text, stats = anonymizer.anonymize(markdown_text, multiplier, entity_types, skip_entities=skip_set)

                anon_filename = f"anon_{os.path.splitext(original_filename)[0]}.md"
                anon_path = TEMP_DIR / anon_filename
                with open(anon_path, "w") as f:
                    f.write(anon_text)

                if os.path.exists(md_path):
                    try:
                        os.unlink(md_path)
                    except Exception:
                        pass

                key_filename = f"{os.path.splitext(original_filename)[0]}.bridgekey.json"
                key_data = anonymizer.create_bridge_key(original_filename, entity_map, multiplier)
                key_path = TEMP_DIR / key_filename
                with open(key_path, "w") as f:
                    json.dump(key_data, f, indent=2)

                total_entities = stats["persons"] + stats["orgs"] + stats.get("locs", 0)
                elapsed = time.time() - start_time
                logger.info("Done: %s (scanned PDF → Markdown) — %d entities (%.2fs)", original_filename, total_entities, elapsed)

                message = f"Anonymized {total_entities} entities"
                return JSONResponse({
                    "success": True,
                    "message": message,
                    "anon_filename": anon_filename,
                    "key_filename": key_filename,
                    "entity_mapping": entity_map,
                    "stats": stats,
                })
            else:
                text, pages = extract_text_from_pdf(file_bytes)
                entity_map, _, stats = anonymizer.anonymize(text, multiplier, entity_types, skip_entities=skip_set)

                counts, pdf_bytes, number_map = anonymize_pdf(file_bytes, entity_map, multiplier)

                anon_filename = f"anon_{original_filename}"
                anon_path = TEMP_DIR / anon_filename
                with open(anon_path, "wb") as f:
                    f.write(pdf_bytes)

                key_filename = f"{os.path.splitext(original_filename)[0]}.bridgekey.json"
                key_data = anonymizer.create_bridge_key(original_filename, entity_map, multiplier)
                key_data["number_mappings"] = number_map
                key_path = TEMP_DIR / key_filename
                with open(key_path, "w") as f:
                    json.dump(key_data, f, indent=2)

                total_entities = stats["persons"] + stats["orgs"] + stats.get("locs", 0)
                elapsed = time.time() - start_time
                logger.info("Done: %s (text PDF) — %d entities, %d numbers (%.2fs)", original_filename, total_entities, counts['numbers'], elapsed)

                message = f"Anonymized {total_entities} entities, {counts['numbers']} numbers"
                return JSONResponse({
                    "success": True,
                    "message": message,
                    "anon_filename": anon_filename,
                    "key_filename": key_filename,
                    "entity_mapping": entity_map,
                    "stats": stats,
                })

        elif file_ext == ".pptx":
            text, prs = extract_text_from_pptx(file_bytes)
            entity_map, _, stats = anonymizer.anonymize(text, multiplier, entity_types, skip_entities=skip_set)

            counts, number_map = anonymize_pptx(prs, entity_map, multiplier)

            anon_filename = f"anon_{original_filename}"
            anon_path = TEMP_DIR / anon_filename
            prs.save(anon_path)

            key_filename = f"{os.path.splitext(original_filename)[0]}.bridgekey.json"
            key_data = anonymizer.create_bridge_key(original_filename, entity_map, multiplier)
            key_data["number_mappings"] = number_map
            key_path = TEMP_DIR / key_filename
            with open(key_path, "w") as f:
                json.dump(key_data, f, indent=2)

            total_entities = stats["persons"] + stats["orgs"] + stats.get("locs", 0)
            elapsed = time.time() - start_time
            logger.info("Done: %s — %d entities, %d numbers (%.2fs)", original_filename, total_entities, counts['numbers'], elapsed)

            message = f"Anonymized {total_entities} entities, {counts['numbers']} numbers"
            return JSONResponse({
                "success": True,
                "message": message,
                "anon_filename": anon_filename,
                "key_filename": key_filename,
                "entity_mapping": entity_map,
                "stats": stats,
            })

        else:
            logger.warning("Unsupported file type: %s", file_ext)
            raise HTTPException(status_code=400, detail=f"Unsupported file type: {file_ext}")

    except MemoryThresholdExceeded:
        logger.warning("Memory threshold exceeded during anonymize: %s", original_filename)
        raise HTTPException(status_code=413, detail="File Too Large — processing exceeded memory limit")
    except Exception as e:
        logger.exception("Anonymize failed: %s (%s)", original_filename, e)
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/download-all/{anon_filename}")
async def download_all(anon_filename: str):
    if not anon_filename.startswith("anon_"):
        raise HTTPException(status_code=400, detail="Invalid filename")

    original = anon_filename[5:]
    key_filename = f"{os.path.splitext(original)[0]}.bridgekey.json"

    anon_path = TEMP_DIR / anon_filename
    key_path = TEMP_DIR / key_filename

    if not anon_path.exists():
        raise HTTPException(status_code=404, detail="Anonymized file not found")
    if not key_path.exists():
        raise HTTPException(status_code=404, detail="Key file not found")

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.write(anon_path, arcname=original)
        zf.write(key_path, arcname=key_filename)

    zip_buffer.seek(0)
    zip_name = f"{os.path.splitext(original)[0]}_anonymized.zip"

    return FileResponse(
        zip_buffer,
        media_type="application/zip",
        filename=zip_name,
        headers={"Content-Disposition": f'attachment; filename="{zip_name}"'}
    )


@app.get("/api/download/{filename}")
async def download_file(filename: str):
    file_path = TEMP_DIR / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found")

    return FileResponse(
        file_path,
        media_type="application/octet-stream",
        filename=filename
    )


@app.post("/api/restore")
async def restore_file(
    file: Optional[UploadFile] = File(None),
    key_file: Optional[UploadFile] = File(None),
    raw_text: Optional[str] = Form(None),
):
    if key_file is None:
        raise HTTPException(status_code=400, detail="Key file is required")

    key_bytes = await key_file.read()
    try:
        key_data = json.loads(key_bytes)
    except json.JSONDecodeError:
        raise HTTPException(status_code=400, detail="Invalid key file format")

    try:
        logger.info("Restore: %s", "raw_text" if raw_text else file.filename if file else "unknown")
        if raw_text:
            restored_text = raw_text
            entity_map = key_data.get("entities", {})
            number_map = key_data.get("number_mappings", {})
            multiplier = key_data.get("multiplier", 1.0)

            for placeholder, original in entity_map.items():
                if placeholder in restored_text:
                    restored_text = restored_text.replace(placeholder, original)

            for orig_str, new_val in number_map.items():
                if str(new_val) in restored_text:
                    restored_text = restored_text.replace(str(new_val), orig_str)

            return JSONResponse({
                "success": True,
                "message": "Text restored successfully",
                "restored_text": restored_text,
            })

        if file is None:
            raise HTTPException(status_code=400, detail="No file or text provided")

        original_filename = file.filename
        file_ext = os.path.splitext(original_filename)[1].lower()
        file_bytes = await file.read()

        if file_ext == ".docx":
            doc = Document(io.BytesIO(file_bytes))
            restore_counts = restore_docx(doc, key_data)

            output_buffer = io.BytesIO()
            doc.save(output_buffer)
            output_buffer.seek(0)

            restored_filename = f"restored_{original_filename}"
            restored_path = TEMP_DIR / restored_filename
            with open(restored_path, "wb") as f:
                f.write(output_buffer.getvalue())

            return JSONResponse({
                "success": True,
                "message": f"Restored {restore_counts['entities']} entities, {restore_counts['numbers']} numbers",
                "restored_filename": restored_filename,
            })

        elif file_ext == ".xlsx":
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
            restore_counts = restore_xlsx(wb, key_data)

            restored_filename = f"restored_{original_filename}"
            restored_path = TEMP_DIR / restored_filename
            wb.save(restored_path)

            return JSONResponse({
                "success": True,
                "message": f"Restored {restore_counts['entities']} entities, {restore_counts['numbers']} numbers",
                "restored_filename": restored_filename,
            })

        elif file_ext == ".pdf":
            restore_counts, pdf_bytes = restore_pdf(file_bytes, key_data)

            restored_filename = f"restored_{original_filename}"
            restored_path = TEMP_DIR / restored_filename
            with open(restored_path, "wb") as f:
                f.write(pdf_bytes)

            return JSONResponse({
                "success": True,
                "message": f"Restored {restore_counts['entities']} entities, {restore_counts['numbers']} numbers",
                "restored_filename": restored_filename,
            })

        elif file_ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(io.BytesIO(file_bytes))
            restore_counts = restore_pptx(prs, key_data)

            restored_filename = f"restored_{original_filename}"
            restored_path = TEMP_DIR / restored_filename
            prs.save(restored_path)

            return JSONResponse({
                "success": True,
                "message": f"Restored {restore_counts['entities']} entities, {restore_counts['numbers']} numbers",
                "restored_filename": restored_filename,
            })

        elif file_ext in (".md", ".txt"):
            restored_text = file_bytes.decode("utf-8")
            entity_map = key_data.get("entities", {})
            number_map = key_data.get("number_mappings", {})

            for placeholder, original in entity_map.items():
                if placeholder in restored_text:
                    restored_text = restored_text.replace(placeholder, original)

            for orig_str, new_val in number_map.items():
                if str(new_val) in restored_text:
                    restored_text = restored_text.replace(str(new_val), orig_str)

            restored_filename = f"restored_{original_filename}"
            restored_path = TEMP_DIR / restored_filename
            with open(restored_path, "w") as f:
                f.write(restored_text)

            entities_replaced = sum(
                1 for p in entity_map if p in file_bytes.decode("utf-8")
            )

            return JSONResponse({
                "success": True,
                "message": f"Restored {entities_replaced} entities, {len(number_map)} numbers",
                "restored_filename": restored_filename,
            })

        else:
            raise HTTPException(status_code=400, detail=f"Unsupported file type: {file_ext}")

    except MemoryThresholdExceeded:
        logger.warning("Memory threshold exceeded during restore")
        raise HTTPException(status_code=413, detail="File Too Large — processing exceeded memory limit")
    except Exception as e:
        logger.exception("Restore failed: %s", e)
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/cleanup")
async def cleanup_temp():
    count = 0
    for file_path in TEMP_DIR.iterdir():
        if file_path.is_file():
            try:
                file_path.unlink()
                count += 1
            except Exception:
                pass
    cleanup_ocr_temp()
    logger.info("Cleanup: removed %d temp files", count)
    return {"success": True}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000)
